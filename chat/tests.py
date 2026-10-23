from unittest.mock import patch
from django.test import TestCase, Client
from chat.backends.agentic_online import (
    calculate as online_calculate,
    get_current_time as online_time,
    get_weather as online_weather,
    reverse_text as online_reverse,
    run_system_command
)
from chat.backends.agentic_offline import (
    calculate as offline_calculate,
    get_current_time as offline_time,
    reverse_text as offline_reverse
)


class GenAIStudioGroupedTests(TestCase):
    def setUp(self):
        self.client = Client()

    # =================================================================
    # GROUP 1: GenAI Chat
    # =================================================================
    def test_group1_online_chat_ui(self):
        res = self.client.get('/chat/')
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "GenAI")

    def test_group1_offline_chat_ui(self):
        res = self.client.get('/chat/offline/')
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "GenAI")

    @patch('chat.views.generate_online_chat')
    def test_group1_online_chat_api(self, mock_chat):
        mock_chat.return_value = {"reply": "Online Gemini reply", "type": "genai_online"}
        res = self.client.get('/chat/api?q=hello')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["reply"], "Online Gemini reply")

    @patch('chat.views.generate_offline_chat')
    def test_group1_offline_chat_api(self, mock_chat):
        mock_chat.return_value = {"reply": "Offline Ollama reply", "type": "genai_offline"}
        res = self.client.get('/chat/offline/api?q=hello&model=llama3.2')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["reply"], "Offline Ollama reply")

    # =================================================================
    # GROUP 2: Agentic AI Online
    # =================================================================
    def test_group2_limited_agent_ui(self):
        res = self.client.get('/chat/agent/limited/')
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Agentic AI Online")

    def test_group2_universal_agent_ui(self):
        res = self.client.get('/chat/agent/universal/')
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Agentic AI Online")

    def test_group2_limited_tools(self):
        self.assertEqual(online_calculate("25 * 4"), "100")
        self.assertTrue(len(online_time()) >= 10)
        self.assertIn("18°C", online_weather("Tokyo"))
        self.assertEqual(online_reverse("hello"), "olleh")

    def test_group2_universal_tool(self):
        with patch('subprocess.run') as mock_sub:
            mock_sub.return_value.stdout = "75"
            mock_sub.return_value.stderr = ""
            mock_sub.return_value.returncode = 0
            self.assertEqual(run_system_command("Get-Battery"), "75")

    @patch('chat.views.run_online_limited_agent')
    def test_group2_limited_agent_api(self, mock_agent):
        mock_agent.return_value = {
            "reply": "It is 12:00",
            "steps": [{"tool": "get_current_time", "output": "12:00"}],
            "type": "agentic_online_limited"
        }
        res = self.client.get('/chat/agent/limited/api?q=what+time')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["reply"], "It is 12:00")

    @patch('chat.views.propose_universal_action')
    def test_group2_universal_agent_propose(self, mock_propose):
        mock_propose.return_value = {
            "needs_confirmation": True,
            "command": "Get-Process",
            "tool": "run_powershell_command",
            "type": "agentic_online_universal"
        }
        res = self.client.get('/chat/agent/universal/api?q=show+processes&confirmed=false')
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.json()["needs_confirmation"])
        self.assertEqual(res.json()["command"], "Get-Process")

    @patch('chat.views.execute_confirmed_command')
    def test_group2_universal_agent_execute_confirmed_password(self, mock_exec):
        mock_exec.return_value = {
            "needs_confirmation": False,
            "reply": "Command output",
            "steps": [{"tool": "run_powershell_command", "output": "done"}],
            "type": "agentic_online_universal"
        }
        res = self.client.get('/chat/agent/universal/api?q=show+processes&confirmed=true&command=Get-Process&password=admin123')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["reply"], "Command output")

    @patch('chat.views.execute_confirmed_command')
    def test_group2_universal_agent_invalid_password_returns_403(self, mock_exec):
        mock_exec.side_effect = PermissionError("Access Denied: Incorrect password")
        res = self.client.get('/chat/agent/universal/api?q=show+processes&confirmed=true&command=Get-Process&password=wrong')
        self.assertEqual(res.status_code, 403)
        self.assertTrue(res.json().get("auth_error"))

    # =================================================================
    # GROUP 3: Agentic AI Offline
    # =================================================================
    def test_group3_offline_agent_ui(self):
        res = self.client.get('/chat/agent/offline/')
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Agentic AI Offline")

    def test_group3_offline_tools(self):
        self.assertEqual(offline_calculate("50 + 50"), "100")
        self.assertTrue(len(offline_time()) >= 10)
        self.assertEqual(offline_reverse("offline"), "enilffo")

    @patch('chat.views.run_offline_agent')
    def test_group3_offline_agent_api(self, mock_agent):
        mock_agent.return_value = {
            "reply": "Calculation is 100",
            "steps": [{"tool": "calculate", "args": {"expression": "50+50"}, "output": "100"}],
            "type": "agentic_offline"
        }
        res = self.client.get('/chat/agent/offline/api?q=calculate+50%2B50&model=llama3.2')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["reply"], "Calculation is 100")

    # =================================================================
    # Backward compatibility routes
    # =================================================================
    def test_backward_compatible_routes(self):
        res1 = self.client.get('/chat/limited-agent/')
        self.assertEqual(res1.status_code, 200)
        res2 = self.client.get('/chat/universal-agent/')
        self.assertEqual(res2.status_code, 200)
        res3 = self.client.get('/chat/agent/')
        self.assertEqual(res3.status_code, 200)

    # =================================================================
    # GROUP 4: Gmail Model Context Protocol (Gmail MCP Studio)
    # =================================================================
    @patch('chat.views.search_emails')
    def test_group4_gmail_mcp_ui(self, mock_search):
        mock_search.return_value = [{"id": "1", "sender": "test@gmail.com", "subject": "Test", "timestamp": "Now"}]
        res = self.client.get('/chat/mcp/gmail/')
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "Gmail MCP Studio")
        self.assertContains(res, "Gmail-MCP-Server")

    def test_group4_gmail_mcp_catalog_api(self):
        res = self.client.get('/chat/mcp/gmail/api?action=catalog')
        self.assertEqual(res.status_code, 200)
        data = res.json()
        tool_names = [t["name"] for t in data["tools"]]
        self.assertIn("search_emails", tool_names)
        self.assertIn("read_email", tool_names)
        self.assertIn("send_email", tool_names)
        self.assertIn("create_draft", tool_names)

    @patch('chat.backends.gmail_mcp_client.call_gmail_tool')
    def test_group4_gmail_mcp_search_and_read(self, mock_call):
        mock_call.return_value = {
            "tool": "search_emails", "output": "Security alert", "is_error": False,
            "rpc_request": {"method": "tools/call"}, "rpc_response": {"result": {}}
        }
        res1 = self.client.get('/chat/mcp/gmail/api?action=call_tool&name=search_emails&args={"query":"security"}')
        self.assertEqual(res1.status_code, 200)
        self.assertIn("Security alert", res1.json()["output"])

    @patch('chat.backends.gmail_mcp_client.call_gmail_tool')
    def test_group4_gmail_mcp_send_and_draft(self, mock_call):
        mock_call.return_value = {
            "tool": "send_email", "output": "Real email sent via Gmail SMTP.", "is_error": False,
            "rpc_request": {"method": "tools/call"}, "rpc_response": {"result": {}}
        }
        res1 = self.client.get('/chat/mcp/gmail/api?action=call_tool&name=send_email&args={"recipient":"test@demo.com","subject":"Test","body":"Hello"}')
        self.assertEqual(res1.status_code, 200)
        self.assertIn("sent", res1.json()["output"].lower())

    @patch('chat.backends.gmail_mcp_client.read_gmail_resource')
    def test_group4_gmail_mcp_read_resource(self, mock_read):
        mock_read.return_value = {"uri": "gmail://inbox/unread", "content": "Real Gmail Unread Count: 5", "rpc_request": {}, "rpc_response": {}}
        res = self.client.get('/chat/mcp/gmail/api?action=read_resource&uri=gmail://inbox/unread')
        self.assertEqual(res.status_code, 200)
        self.assertIn("Real Gmail", res.json()["content"])

    @patch('chat.backends.gmail_mcp_client.run_gmail_agent')
    def test_group4_gmail_mcp_agent_query(self, mock_agent):
        mock_agent.return_value = {
            "reply": "Found security emails", "steps": [{"tool": "search_emails"}], "type": "gmail_mcp_agent"
        }
        res = self.client.get('/chat/mcp/gmail/api?q=search+emails')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["type"], "gmail_mcp_agent")

    # =================================================================
    # GROUP 5: RAG Studio (Retrieval-Augmented Generation)
    # =================================================================
    def test_group5_rag_ui(self):
        res = self.client.get('/chat/rag/')
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "RAG Studio")
        self.assertContains(res, "Retrieval-Augmented Generation")

    def test_group5_rag_direct_shortcut_ui(self):
        res = self.client.get('/rag/')
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, "RAG Studio")

    def test_group5_rag_documents_api(self):
        # Reset defaults so resume.pdf is loaded
        from chat.backends.rag_service import kb
        kb.reset_defaults()
        res = self.client.get('/chat/rag/api?action=documents')
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertIn("documents", data)
        self.assertGreaterEqual(data["total_docs"], 1)
        self.assertTrue(any("Resume" in d["title"] for d in data["documents"]))

    def test_group5_rag_add_delete_and_clear_documents(self):
        # Clear KB first to test add/delete isolation
        self.client.post('/chat/rag/api', data={"action": "clear"})

        # Ingest custom document
        res_add = self.client.post('/chat/rag/api', data={
            "action": "add_document",
            "title": "Quantum Computing FAQ",
            "content": "Quantum supremacy refers to quantum computers solving problems beyond classical computers."
        })
        self.assertEqual(res_add.status_code, 200)
        self.assertEqual(res_add.json()["status"], "ok")
        doc_id = res_add.json()["document"]["id"]

        # Verify it was added
        res_docs = self.client.get('/chat/rag/api?action=documents')
        titles = [d["title"] for d in res_docs.json()["documents"]]
        self.assertIn("Quantum Computing FAQ", titles)

        # Delete document
        res_del = self.client.get(f'/chat/rag/api?action=delete_document&doc_id={doc_id}')
        self.assertEqual(res_del.status_code, 200)
        self.assertEqual(res_del.json()["status"], "ok")
        self.assertEqual(res_del.json()["total_docs"], 0)

        # Clear/Reset KB
        res_clear = self.client.post('/chat/rag/api', data={"action": "clear"})
        self.assertEqual(res_clear.status_code, 200)
        self.assertEqual(res_clear.json()["status"], "ok")
        self.assertEqual(res_clear.json()["total_docs"], 0)

        # Restore defaults
        from chat.backends.rag_service import kb
        kb.reset_defaults()

    def test_group5_rag_file_upload(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        # Upload a text document file
        uploaded_txt = SimpleUploadedFile("cloud_architecture.txt", b"Cloud microservices communicate via REST and gRPC.")
        res = self.client.post('/chat/rag/api', data={
            "action": "add_document",
            "category": "Engineering",
            "file": uploaded_txt
        })
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["status"], "ok")
        self.assertIn("Cloud Architecture", res.json()["document"]["title"])

        # Clean up
        self.client.post('/chat/rag/api', data={"action": "clear"})
        from chat.backends.rag_service import kb
        kb.reset_defaults()

    def test_group5_rag_docx_upload(self):
        import io
        import docx
        from django.core.files.uploadedfile import SimpleUploadedFile

        # Create in-memory docx
        doc = docx.Document()
        doc.add_heading('Technical Specifications', level=1)
        doc.add_paragraph('System requires Python 3.11+, Django 4.2, and PostgreSQL.')
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)

        uploaded_docx = SimpleUploadedFile(
            "technical_specs.docx",
            buf.read(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        res = self.client.post('/chat/rag/api', data={
            "action": "add_document",
            "category": "Engineering",
            "file": uploaded_docx
        })
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["status"], "ok")
        self.assertIn("Technical Specs", res.json()["document"]["title"])

        # Clean up
        self.client.post('/chat/rag/api', data={"action": "clear"})
        from chat.backends.rag_service import kb
        kb.reset_defaults()

    @patch('chat.views.execute_rag_pipeline')
    def test_group5_rag_query_api(self, mock_rag):
        mock_rag.return_value = {
            "reply": "RAG chunks text into manageable segments for vector indexing [1].",
            "retrieved_chunks": [
                {
                    "rank": 1,
                    "doc_id": "doc_rag_fundamentals",
                    "doc_title": "RAG Architecture",
                    "chunk_id": "doc_rag_fundamentals_chunk_0",
                    "text": "Chunking splits long documents into smaller semantic windows.",
                    "similarity": 0.88,
                    "method": "embedding_cosine"
                }
            ],
            "stats": {
                "top_k": 3,
                "retrieved_count": 1,
                "similarity_mode": "gemini-embedding-001 (cosine)",
                "generation_mode": "gemini-2.5-flash",
                "retrieval_ms": 12.5,
                "generation_ms": 150.2
            },
            "prompt_injected": "System Prompt with Grounded Context"
        }
        res = self.client.get('/chat/rag/api?q=What+is+chunking&k=3')
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertIn("reply", data)
        self.assertEqual(len(data["retrieved_chunks"]), 1)
        self.assertIn("[1]", data["reply"])
        self.assertEqual(data["stats"]["similarity_mode"], "gemini-embedding-001 (cosine)")


